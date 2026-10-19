"""Lit VOXITY - ASSETS.xlsx, calcule l'équilibrage de chaque tuile, écrit l'xlsx enrichi + data.json."""
import openpyxl, json, hashlib, sys
src, out_xlsx, out_json = sys.argv[1:4]
wb = openpyxl.load_workbook(src)
ws = wb['Voxity']
hdr = [c.value for c in ws[1]]
rows = [r for r in ws.iter_rows(min_row=2) if r[0].value]

CATS = {'Houses':0,'Shops':1,'Fun':2,'Production':3,'Services':4,'Roads':5}
RAR = {'STARTER':0,'⭐':1,'⭐⭐':2,'⭐⭐⭐':3,'⭐⭐⭐⭐':4}
RANGE = [1,2,3,4,99]
CM = [1,2.2,5,11,26]            # multiplicateur de coût par rareté
MH = [1,1.8,3.2,5.5,9]          # multiplicateur de capacité d'habitants
def h(s, salt=''):
    return int(hashlib.md5((s+salt).encode()).hexdigest()[:8],16)/0xffffffff
def jit(s, salt, amp=0.12): return 1 + (h(s,salt)*2-1)*amp
def r5(x): return max(5, int(round(x/5.0))*5)
def pct(x): return round(x)/100.0

# (coût de base, valeur de base) par sous-catégorie
HOUSE = {'Maison':(45,5),'Résidence':(110,12),'Batiment':(80,9),'Villa':(150,14),'Gratte Ciel':(400,36)}
SHOP  = {'Amplette':60,'Mobilier':40,'Nourriture':90,'Restaurant':140,'Esthetique':160}
FUN   = {'Divertissement':220,'Montagne':160,'Nature':70,'Plage':90,'Sport':130}
PROD  = {'Agriculture':(80,.9,.5),'Chantier':(100,.8,.8),'Energie':(220,1.2,1.2),'Fabrication':(180,1.3,1.4),'Technologie':(300,1.4,1.0)}
SERV  = {'Education':160,'Institution':220,'Public':140,'Santé':200}
ROAD  = {'Chemin':8,'Route':12,'Autoroute':30,'Cyclisme':15,'Ferroviaire':40,'Maritime':25,'Aviation':120,'Pont':20}
NETWORK = {'Chemin','Route','Autoroute','Cyclisme','Ferroviaire','Pont'}

data=[]; themes=[]
for r in rows:
    v=[c.value for c in r]
    model,name,cat,sub,pop,nowalk,clock,rlock,theme,rar = v[:10]
    s = RAR[rar]; subs=[x.strip() for x in sub.split(',')]; p=subs[0]
    if theme not in themes: themes.append(theme)
    k=model
    cost=hab=money=capH=capM=mH=mM=bM=bH=vote=0
    if cat=='Houses':
        c0,C0 = HOUSE[p]; capH=max(2,round(C0*MH[s]*jit(k,'c')/3)); cost=c0*CM[s]*jit(k,'k'); hab=round(capH*.65)
        bM=pct(5+3*s)
        vote=0
    elif cat=='Shops':
        c0=SHOP[p]; cost=r5(c0*CM[s]*jit(k,'k')); money=r5(cost*[1.10,1.05,.95,.90,.85][s]*jit(k,'m',.03)); capM=int(round(cost*4/50.0))*50
        hab=round(cost/100); vote=0
    elif cat=='Fun':
        c0=FUN[p]; cost=r5(c0*CM[s]*jit(k,'k')); money=r5(cost*[1.0,.95,.9,.85,.8][s]*jit(k,'m',.03)); hab=round(cost/(54+12*s)*jit(k,'h')); capM=0
        vote=1+s
    elif cat=='Production':
        c0,wb_,wm=PROD[p]; cost=r5(c0*CM[s]*jit(k,'k')); money=r5(cost*.5)
        bM=pct((30+6*s)*wb_*jit(k,'b',.08)); mH=pct((4+3*s)*wm*jit(k,'w',.08))
    elif cat=='Services':
        c0=SERV[p]; cost=c0*CM[s]*jit(k,'k')
        bH=pct((6+4*s)*jit(k,'b',.08)); mM=pct(round(bH*100)*.5); vote=2+2*s
    else:
        if name.lower().startswith('hill') or theme=='#HILL#': c0=5
        elif len(subs)>=4 and 'Pont' not in subs: c0=20       # véhicules
        elif p in ROAD: c0=ROAD[p]
        else: c0=6                                              # Elements
        cost=r5(c0*CM[s]*jit(k,'k'))
        if p in NETWORK and theme!='#HILL#':
            bM=bH=pct(.6*(1+s))
    cost=r5(cost); money=r5(money) if money else 0
    data.append(dict(model=model,name=name,cat=cat,sub=sub,theme=theme,rar=rar,s=s,cost=cost,hab=hab,money=money,
        capH=capH,capM=capM,mH=mH,mM=mM,bM=bM,bH=bH,rng=RANGE[s],vote=vote,rlock=rlock,subs=subs,nowalk=bool(nowalk)))
    # écriture dans l'Excel
    vals={'Coût':cost,'Habitants rapportés':hab or None,'Argent rapporté':money or None,"Limite d'habitant augmenté":capH or None,
          "Limite d'argent augmenté":capM or None,'Malus Habitant':mH or None,'Malus Argent':mM or None,
          "Bonus d'argent rapporté":bM or None,"Bonus d'habitant rapporté":bH or None,"Portée d'effet":RANGE[s] if RANGE[s]<99 else 'Carte','Pts de vote':vote or None}
    for key,val in vals.items(): r[hdr.index(key)].value=val
wb.save(out_xlsx)

best={}
for i,d_ in enumerate(data):
    b=best.get(d_['name'])
    if b is None or d_['s']<data[b]['s']: best[d_['name']]=i
keep=set(best.values())
col=hdr.index('Asset Rarity')+ (len(hdr)-hdr.index('Asset Rarity'))
ws.cell(row=1,column=len(hdr)+1,value='Tuile retenue (1 par Asset Name)')
for i,r in enumerate(rows): ws.cell(row=r[0].row,column=len(hdr)+1,value=(i in keep))
wb.save(out_xlsx)
data=[d_ for i,d_ in enumerate(data) if i in keep]
themes=[]
for d_ in data:
    if d_['theme']!='#HILL#' and d_['theme'] not in themes: themes.append(d_['theme'])
cnt={t:sum(1 for d in data if d['theme']==t) for t in themes}
# bonus de thème (additif) : type et valeur déterministes
KINDS=['attr','rend','pop','eco','porte']
tb=[]
for i,t in enumerate(themes):
    n=cnt[t]; kind=KINDS[i%4]
    if n>=25 and h(t,'p')<.45: kind='porte'
    val={'attr':pct(2+.03*min(n,100)),'rend':pct(3+.04*min(n,100)),'pop':pct(3+.05*min(n,100)),
         'eco':pct(2+.03*min(n,100)),'porte':1}[kind]
    tb.append(dict(theme=t,n=n,kind=kind,val=val,ecoCat=int(h(t,'e')*5) if kind=='eco' else None))
cats=['Habitations','Commerces','Loisirs','Productions','Services','Routes']
d=dict(cats=cats,themes=themes,themeBonus=tb,
  tiles=[[x['model'],x['name'],CATS[x['cat']],x['sub'],themes.index(x['theme']) if x['theme']!='#HILL#' else -1,x['s'],x['cost'],x['hab'],x['money'],
   x['capH'],x['capM'],x['mH'],x['mM'],x['bM'],x['bH'],x['vote'],x['rlock'],0 if x['nowalk'] else 1] for x in data])
json.dump(d,open(out_json,'w'),ensure_ascii=False,separators=(',',':'))
print(len(data),len(themes))
import collections
for c in cats: pass
